"""
Excelファイルの全シートを読み込み、構造とデータをJSON形式で出力するスクリプト。
使い方: python read_excel.py <Excelファイルのパス>
"""
import sys
import json
import openpyxl

def read_excel(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    result = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            # 全セルが空の行はスキップ
            if any(cell is not None for cell in row):
                rows.append(list(row))
        result[sheet_name] = rows

    return result

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("使い方: python read_excel.py <Excelファイルのパス>")
        sys.exit(1)

    filepath = sys.argv[1]
    data = read_excel(filepath)

    print(json.dumps(data, ensure_ascii=False, default=str))
